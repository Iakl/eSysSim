import os
import shutil
import argparse
import openpyxl
from datetime import datetime


def convert_params(src_path: str, dst_path: str) -> None:
    src_wb = openpyxl.load_workbook(src_path, data_only=True)
    dst_wb = openpyxl.Workbook()

    bus_alias_to_idx = {}

    for sheet_name in src_wb.sheetnames:
        ws = src_wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [h for h in rows[0]]
        data = [r for r in rows[1:] if any(v is not None for v in r)]
        if not data:
            continue

        if sheet_name == "Bus":
            dst_ws = dst_wb.active
            dst_ws.title = "Bus"
            dst_ws.append(["index", "alias", "vn_kv", "v_pu_min", "v_pu_max"])
            for r in data:
                idx = r[0]
                alias = r[1]
                bus_alias_to_idx[alias] = idx
                dst_ws.append([idx, alias, r[2], r[3], r[4]])

        elif sheet_name == "ElectricGrid":
            dst_ws = dst_wb.create_sheet("ElectricGrid")
            dst_ws.append([
                "index", "alias", "connection1", "is_active", "v_set_pu",
                "p_mw_term", "p_min_mw", "p_max_mw", "q_max_mvar", "q_min_mvar",
                "purchase_price", "sell_price", "p_term_mode", "tariff_zone",
                "tariff", "purchase_price_mode"
            ])
            for r in data:
                conn_idx = bus_alias_to_idx.get(r[2], 0)
                dst_ws.append([
                    r[0], r[1], conn_idx, r[3], r[4],
                    r[5], r[6], r[7], r[8], r[9],
                    r[10], r[11], r[12], r[13],
                    r[14], r[15]
                ])

        elif sheet_name == "ElectricLoad":
            dst_ws = dst_wb.create_sheet("ElectricLoad")
            dst_ws.append([
                "index", "alias", "connection1", "load_type", "is_active",
                "controllable", "priority", "p_mw", "q_mvar", "p_mw_mode",
                "q_mvar_mode", "p_max_mw", "p_min_mw", "q_max_mvar", "q_min_mvar"
            ])
            for r in data:
                conn_idx = bus_alias_to_idx.get(r[2], 0)
                p_mode = r[5] if len(r) > 5 else "profile"
                p_mw = r[6] if len(r) > 6 else 0
                is_active = r[3] if len(r) > 3 else 1
                controllable = r[4] if len(r) > 4 else 0
                dst_ws.append([
                    r[0], r[1], conn_idx, "critical", is_active,
                    controllable, 0,
                    p_mw, 0, p_mode, "fixed",
                    None, None, None, None
                ])

        elif sheet_name == "Line":
            dst_ws = dst_wb.create_sheet("Line")
            dst_ws.append([
                "index", "alias", "from_bus", "to_bus", "is_active",
                "p_mw_max", "r_pu", "x_pu", "ang_grad_min", "ang_grad_max",
                "overload_pu"
            ])
            for r in data:
                from_bus = bus_alias_to_idx.get(r[2], 0)
                to_bus = bus_alias_to_idx.get(r[3], 0)
                s_mva_max = r[5] if len(r) > 5 else 300
                r_ohm = r[6] if len(r) > 6 else 0
                x_ohm = r[7] if len(r) > 7 else 0
                vn_kv = 230
                z_base = vn_kv ** 2 / s_mva_max
                r_pu = r_ohm / z_base
                x_pu = x_ohm / z_base
                overload = r[9] if len(r) > 9 else 0.8
                dst_ws.append([
                    r[0], r[1], from_bus, to_bus, r[4] if len(r) > 4 else 1,
                    s_mva_max, r_pu, x_pu, -180, 180, overload
                ])

        elif sheet_name == "Battery":
            dst_ws = dst_wb.create_sheet("Battery")
            dst_ws.append([
                "index", "alias", "connection1", "is_active", "p_mw", "q_mvar",
                "p_min_mw", "p_max_mw", "q_min_mvar", "q_max_mvar",
                "c_degr", "capacity_mwh", "soc_ini_pu", "soc_max_pu", "soc_min_pu",
                "eff_ch", "eff_dch", "controllable", "priority"
            ])
            for r in data:
                conn_idx = bus_alias_to_idx.get(r[2], 0)
                dst_ws.append([
                    r[0], r[1], conn_idx, r[3], r[4], r[5],
                    r[6], r[7], r[8], r[9],
                    r[10], r[11], r[12], r[13], r[14],
                    r[15], r[16], r[17], r[18] if len(r) > 18 else 1
                ])

        elif sheet_name == "Generator":
            dst_ws = dst_wb.create_sheet("Generator")
            dst_ws.append([
                "index", "alias", "connection1", "p_mw", "q_mvar",
                "p_min_mw", "p_max_mw", "q_min_mvar", "q_max_mvar",
                "p_mw_ini", "p_c2", "p_c1", "is_active", "controllable",
                "v_set_pu", "priority", "p_max_mw_mode", "d_mw_max"
            ])
            for r in data:
                conn_idx = bus_alias_to_idx.get(r[2], 0)
                p_max_mode = r[18] if len(r) > 18 else "fixed"
                d_mw_max = r[17] if len(r) > 17 else 6000
                dst_ws.append([
                    r[0], r[1], conn_idx, r[4], r[5],
                    r[6], r[7], r[8], r[9],
                    r[3], r[12], r[13], r[10], r[16],
                    r[11], 1, p_max_mode, d_mw_max
                ])

    dst_wb.save(dst_path)
    print(f"  params.xlsx: {list(dst_wb.sheetnames)}")


def convert_profiles(src_path: str, dst_path: str) -> None:
    src_wb = openpyxl.load_workbook(src_path, data_only=True)
    dst_wb = openpyxl.Workbook()

    for sheet_name in src_wb.sheetnames:
        ws = src_wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        src_headers = [h for h in rows[0]]
        data = rows[1:]

        if sheet_name == "PCC":
            dst_ws = dst_wb.active
            dst_ws.title = "grid0"
            dst_ws.append(["date", "purchase_price", "sell_price"])
            for r in data:
                dt = r[0]
                if isinstance(dt, datetime):
                    dt_str = dt.strftime("%d-%b-%Y %H:%M")
                else:
                    dt_str = str(dt)
                price = r[1] if r[1] is not None else 0
                dst_ws.append([dt_str, price, price * 0.35])

        elif sheet_name == "PV":
            dst_ws = dst_wb.create_sheet("PV")
            dst_ws.append(["date", "p_max_mw", "q_max_mvar"])
            for r in data:
                dt = r[0]
                if isinstance(dt, datetime):
                    dt_str = dt.strftime("%d-%b-%Y %H:%M")
                else:
                    dt_str = str(dt)
                p_max = r[1] if r[1] is not None else 0
                dst_ws.append([dt_str, p_max, 0])

        else:
            dst_ws = dst_wb.create_sheet(sheet_name)
            dst_ws.append(["date", "p_mw", "q_mvar"])
            for r in data:
                dt = r[0]
                if isinstance(dt, datetime):
                    dt_str = dt.strftime("%d-%b-%Y %H:%M")
                else:
                    dt_str = str(dt)
                p_mw = r[1] if r[1] is not None else 0
                q_mvar = r[2] if len(r) > 2 and r[2] is not None else 0
                dst_ws.append([dt_str, p_mw, q_mvar])

    dst_wb.save(dst_path)
    print(f"  profile.xlsx: {list(dst_wb.sheetnames)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert datacenter scenarios")
    parser.add_argument("--src", required=True, help="Source scenarios directory")
    parser.add_argument("--dst", required=True, help="Destination scenarios directory")
    parser.add_argument(
        "--scenarios", nargs="+", default=["scenario1:datacenter_1", "scenario2:datacenter_2"],
        help="Scenario mappings as src_name:dst_name (default: scenario1:datacenter_1 scenario2:datacenter_2)"
    )
    args = parser.parse_args()

    for mapping in args.scenarios:
        src_name, dst_name = mapping.split(":")
        src_dir = os.path.join(args.src, src_name)
        dst_dir = os.path.join(args.dst, dst_name)
        print(f"\n=== Converting {src_name} -> {dst_name} ===")
        convert_params(
            os.path.join(src_dir, "params.xlsx"),
            os.path.join(dst_dir, "params.xlsx")
        )
        convert_profiles(
            os.path.join(src_dir, "profile.xlsx"),
            os.path.join(dst_dir, "profiles.xlsx")
        )
