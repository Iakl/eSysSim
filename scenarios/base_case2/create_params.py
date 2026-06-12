import openpyxl


def create_params_xlsx(output_path: str) -> None:
    wb = openpyxl.Workbook()

    ws_bus = wb.active
    ws_bus.title = "Bus"
    ws_bus.append(["index", "alias", "vn_kv", "v_pu_min", "v_pu_max"])
    ws_bus.append([0, "bus0", 6, 0.9, 1.1])
    ws_bus.append([1, "bus1", 6, 0.9, 1.1])

    ws_grid = wb.create_sheet("ElectricGrid")
    ws_grid.append([
        "index", "alias", "connection1", "is_active", "v_set_pu",
        "p_mw_term", "p_min_mw", "p_max_mw", "q_max_mvar", "q_min_mvar",
        "purchase_price", "sell_price", "p_term_mode", "tariff_zone",
        "tariff", "purchase_price_mode"
    ])
    ws_grid.append([
        0, "grid0", 0, 1, 1.0, 0.8, 0, 1.0, 1.0, -1.0,
        6, 2, "fixed", "peninsula", "6_1", "profile"
    ])

    ws_load = wb.create_sheet("ElectricLoad")
    ws_load.append([
        "index", "alias", "connection1", "load_type", "is_active",
        "controllable", "priority", "p_mw", "q_mvar", "p_mw_mode",
        "q_mvar_mode", "p_max_mw", "p_min_mw", "q_max_mvar", "q_min_mvar"
    ])
    ws_load.append([
        0, "CL1", 1, "critical", 1, 0, 0, -1.0, -0.6, "profile",
        "profile", None, None, None, None
    ])

    ws_line = wb.create_sheet("Line")
    ws_line.append([
        "index", "alias", "from_bus", "to_bus", "is_active",
        "p_mw_max", "r_pu", "x_pu", "ang_grad_min", "ang_grad_max",
        "overload_pu"
    ])
    ws_line.append([
        0, "line0", 0, 1, True, 1.0, 0.001, 0.01, -1.0, 1.0, 1.0
    ])

    wb.save(output_path)


if __name__ == "__main__":
    import os
    script_dir = os.path.dirname(os.path.abspath(__file__))
    output = os.path.join(script_dir, "params.xlsx")
    create_params_xlsx(output)
    print(f"Params saved to {output}")
