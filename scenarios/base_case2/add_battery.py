import openpyxl
import os

scenario_dir = os.path.dirname(os.path.abspath(__file__))

wb = openpyxl.load_workbook(os.path.join(scenario_dir, "params.xlsx"))

ws = wb.create_sheet("Battery")
ws.append([
    "index", "alias", "connection1", "is_active", "p_mw", "q_mvar",
    "p_min_mw", "p_max_mw", "q_min_mvar", "q_max_mvar",
    "c_degr", "capacity_mwh", "soc_ini_pu", "soc_max_pu", "soc_min_pu",
    "eff_ch", "eff_dch", "controllable", "priority"
])
ws.append([
    0, "bat0", 1, 1, 0, 0,
    -0.5, 0.5, -0.5, 0.5,
    0.13, 1.0, 0.6, 0.8, 0.2,
    0.95, 0.95, True, 1
])

wb.save(os.path.join(scenario_dir, "params.xlsx"))
print("Battery added to params.xlsx")

wb_profiles = openpyxl.load_workbook(os.path.join(scenario_dir, "profiles.xlsx"))

ws_bat = wb_profiles.create_sheet("bat0")
ws_bat.append(["date", "p_max_mw", "q_max_mvar"])

from datetime import datetime, timedelta
start = datetime(2017, 1, 1, 0, 0)
for i in range(192):
    dt = start + timedelta(minutes=15 * i)
    ws_bat.append([dt.strftime("%d-%b-%Y %H:%M"), 0.5, 0.5])

wb_profiles.save(os.path.join(scenario_dir, "profiles.xlsx"))
print("bat0 profile added to profiles.xlsx")
